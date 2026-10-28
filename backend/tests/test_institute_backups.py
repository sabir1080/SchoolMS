import datetime
import gzip
import io
import json
import re

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.accounts.models import Role, User
from apps.backups import institute
from apps.backups.models import InstituteSnapshot
from apps.backups.services import create_backup
from apps.fees.models import Invoice, Payment
from apps.messaging.models import MessagingSettings
from apps.payments.models import GatewayAccount
from apps.students.models import Student
from apps.tenants.demo import seed_demo
from apps.tenants.models import Tenant

TODAY = datetime.date(2026, 10, 14)


def rows_of(tenant):
    """Everything that belongs to an institute, comparable before/after (times to the millisecond: the full
    platform backup keeps milliseconds only)."""
    return sorted(re.sub(r"(\.\d{3})\d{3}", r"\1", json.dumps(r, sort_keys=True, default=str))
                  for r in institute.export_rows(tenant, include_secrets=True))


def setup_three(make_tenant, make_user):
    demo = seed_demo(today=TODAY)
    GatewayAccount.objects.create(tenant=demo, provider="jazzcash", merchant_id="MC1", password="pw", secret="salt", is_enabled=True)
    MessagingSettings.objects.update_or_create(tenant=demo, defaults={"twilio_auth_token": "tok-123"})
    others = []
    for slug in ("beta", "gamma"):
        t = make_tenant(slug)
        make_user("admin", Role.ADMIN, t)  # same username as the demo admin
        for i in range(3):
            Student.objects.create(tenant=t, admission_no=str(i), first_name=f"{slug}{i}", gender="male")
        others.append(t)
    return demo, others


def upload(content, name="backup.json.gz"):
    return SimpleUploadedFile(name, content, content_type="application/gzip")


def test_restore_one_institute_leaves_the_others_untouched(api, super_admin, make_tenant, make_user):
    demo, others = setup_three(make_tenant, make_user)
    before = {t.slug: rows_of(t) for t in [demo, *others]}
    file = institute.export_file(demo, include_secrets=False)  # what the school downloads
    assert b"tok-123" not in gzip.decompress(file) and b'"salt"' not in gzip.decompress(file)

    # The demo school's data goes wrong: students deleted, a bill changed, a new student added.
    Student.objects.filter(tenant=demo, admission_no__in=["1", "2"]).update(status="left")
    inv = Invoice.objects.filter(tenant=demo).first()
    Invoice.objects.filter(pk=inv.pk).update(notes="wrong")
    Student.objects.create(tenant=demo, admission_no="X-99", first_name="Late", gender="male")
    Payment.objects.filter(tenant=demo).first().delete()

    api.force_authenticate(super_admin)
    url = f"/api/platform/tenants/{demo.pk}/restore/"
    preview = api.post(url, {"file": upload(file)}, format="multipart").json()
    assert preview["restored"] is False and preview["source"] == "institute"
    assert preview["backup"]["Students"] == 32 and preview["current"]["Students"] == 33
    assert "confirm_code" in api.post(url, {"file": upload(file), "confirm_code": "beta"}, format="multipart").json()

    res = api.post(url, {"file": upload(file), "confirm_code": "DEMO"}, format="multipart").json()
    assert res["restored"] is True and res["current"]["Students"] == 32
    assert rows_of(demo) == before["demo"]  # exactly as at backup time, secrets kept from before
    for t in others:
        assert rows_of(t) == before[t.slug]  # untouched
    assert GatewayAccount.objects.get(tenant=demo).secret == "salt"
    assert MessagingSettings.objects.get(tenant=demo).twilio_auth_token == "tok-123"
    assert User.objects.get(tenant=demo, username="admin").check_password("demo1234")
    assert not Student.objects.filter(tenant=demo, admission_no="X-99").exists()
    # New rows after the restore still get fresh ids.
    s = Student.objects.create(tenant=demo, admission_no="NEW", first_name="New", gender="male")
    assert s.pk > max(Student.objects.exclude(pk=s.pk).values_list("pk", flat=True))

    # The data from just before the restore was saved and can be put back (undo).
    snap = InstituteSnapshot.objects.get(tenant=demo)
    assert snap.reason == "before restore"
    listed = api.get(f"/api/platform/tenants/{demo.pk}/snapshots/").json()
    assert listed[0]["id"] == snap.pk
    undo = api.post(f"/api/platform/tenants/{demo.pk}/snapshots/{snap.pk}/", {"confirm_code": "demo"}, format="json").json()
    assert undo["restored"] is True
    assert Student.objects.filter(tenant=demo, admission_no="X-99").exists()
    for t in others:
        assert rows_of(t) == before[t.slug]


def test_restore_one_institute_from_the_full_backup(api, super_admin, make_tenant, make_user, tmp_path):
    demo, others = setup_three(make_tenant, make_user)
    beta = others[0]
    before = {t.slug: rows_of(t) for t in [demo, *others]}
    path = create_backup(str(tmp_path))
    Student.objects.filter(tenant=beta).delete()
    Student.objects.filter(tenant=demo, admission_no="1").update(first_name="Changed after backup")

    api.force_authenticate(super_admin)
    with open(path, "rb") as fh:
        content = fh.read()
    preview = api.post(f"/api/platform/tenants/{beta.pk}/restore/", {"file": upload(content)}, format="multipart").json()
    assert preview["source"] == "full" and preview["backup"]["Students"] == 3 and "Students" not in preview["current"]
    api.post(f"/api/platform/tenants/{beta.pk}/restore/", {"file": upload(content), "confirm_code": "beta"}, format="multipart")
    assert rows_of(beta) == before["beta"]
    # The demo school's later change is kept: only beta went back in time.
    assert Student.objects.get(tenant=demo, admission_no="1").first_name == "Changed after backup"
    assert rows_of(others[1]) == before["gamma"]


def test_wrong_files_are_refused(api, super_admin, make_tenant):
    alpha, beta = make_tenant("alpha"), make_tenant("beta")
    api.force_authenticate(super_admin)
    url = f"/api/platform/tenants/{alpha.pk}/restore/"
    other = institute.export_file(beta, include_secrets=False)
    assert "for institute “beta”" in str(api.post(url, {"file": upload(other)}, format="multipart").json())
    assert "not a SchoolMS backup" in str(api.post(url, {"file": upload(b"hello")}, format="multipart").json())
    assert "Choose a backup file" in str(api.post(url, {}, format="multipart").json())


def test_school_downloads_its_own_backup_and_excel(api, admin_a, admin_b, school_a, super_admin):
    Student.objects.create(tenant=school_a, admission_no="7", first_name="Zara", gender="female")
    api.force_authenticate(admin_a)
    assert api.get("/api/institute/backup/status/").json()["last_file"] is None
    res = api.get("/api/institute/backup/")
    assert res.status_code == 200 and res["Content-Disposition"].startswith('attachment; filename="alpha-backup-')
    payload = json.loads(gzip.decompress(res.content))
    assert payload["institute"] == "alpha" and payload["with_secrets"] is False
    assert {r["fields"].get("first_name") for r in payload["rows"] if r["model"] == "students.student"} == {"Zara"}
    assert not any(r["model"] in ("subscriptions.subscription", "agreements.signedagreement") for r in payload["rows"])
    status = api.get("/api/institute/backup/status/").json()
    assert status["days_since_file"] == 0

    xlsx = api.get("/api/institute/backup/excel/")
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Students", "Bills", "Payments", "Results", "Attendance (this session)", "Staff"]
    assert wb["Students"]["B5"].value == "Zara"

    # Other institutes' admins see only their own; school admins cannot use the owner's restore.
    api.force_authenticate(admin_b)
    other = json.loads(gzip.decompress(api.get("/api/institute/backup/").content))
    assert other["institute"] == "beta"
    assert api.get(f"/api/platform/tenants/{school_a.pk}/backup/").status_code == 403
    assert api.post(f"/api/platform/tenants/{school_a.pk}/restore/", {}, format="multipart").status_code == 403

    # The owner's copy keeps the payment keys (for a complete restore).
    GatewayAccount.objects.create(tenant=school_a, provider="easypaisa", merchant_id="1", secret="0123456789abcdef")
    api.force_authenticate(super_admin)
    owner = json.loads(gzip.decompress(api.get(f"/api/platform/tenants/{school_a.pk}/backup/").content))
    assert owner["with_secrets"] is True and "0123456789abcdef" in json.dumps(owner)


def test_expired_school_can_still_download(api, admin_a, school_a):
    from django.utils import timezone

    sub = school_a.subscription
    sub.current_period_end = timezone.now() - datetime.timedelta(days=3)
    sub.save()
    api.force_authenticate(admin_a)
    assert api.get("/api/institute/backup/").status_code == 200
    assert Tenant.objects.count() == 1


def test_school_admin_restores_its_own_signed_file(api, admin_a, school_a, make_tenant, make_user, super_admin):
    from apps.messaging.models import PushSubscription
    from apps.recyclebin.models import DeletedItem

    teacher = make_user("teacher_a", Role.TEACHER, school_a)
    beta = make_tenant("beta")
    Student.objects.create(tenant=beta, admission_no="1", first_name="Bilal", gender="male")
    beta_before = rows_of(beta)
    Student.objects.create(tenant=school_a, admission_no="1", first_name="Zara", gender="female")
    GatewayAccount.objects.create(tenant=school_a, provider="jazzcash", merchant_id="M", password="pw", secret="salt")
    push = PushSubscription.objects.create(tenant=school_a, user=admin_a, endpoint="https://fcm.googleapis.com/x", p256dh="k", auth="a")

    api.force_authenticate(admin_a)
    file = api.get("/api/institute/backup/").content
    late = api.post("/api/students/", {"first_name": "Late", "gender": "male"}, format="json").json()["id"]
    zara = Student.objects.get(first_name="Zara")
    assert api.delete(f"/api/students/{zara.pk}/").status_code == 204
    bin_item = DeletedItem.objects.get(tenant=school_a)

    url = "/api/institute/backup/restore/"
    preview = api.post(url, {"file": upload(file)}, format="multipart").json()
    assert preview["restored"] is False and preview["signed"] is True
    assert preview["backup"]["Students"] == 1 and preview["current"]["Students"] == 1
    assert "confirm" in str(api.post(url, {"file": upload(file), "confirm_code": "beta"}, format="multipart").json())
    res = api.post(url, {"file": upload(file), "confirm_code": "ALPHA"}, format="multipart").json()
    assert res["restored"] is True
    assert set(Student.objects.filter(tenant=school_a).values_list("first_name", flat=True)) == {"Zara"}
    assert not Student.objects.filter(pk=late).exists()
    assert rows_of(beta) == beta_before
    # Kept: payment keys (blank in the school's file), the admin's phone for notifications, who deleted what.
    assert GatewayAccount.objects.get(tenant=school_a).secret == "salt"
    assert PushSubscription.objects.filter(pk=push.pk, user=admin_a).exists()
    bin_item.refresh_from_db()
    assert bin_item.deleted_by_id == admin_a.pk
    snap = InstituteSnapshot.objects.get(tenant=school_a)
    assert snap.created_by_id == admin_a.pk

    # Saved copies: listed and put back, never downloaded by the school (they hold payment keys).
    assert [s["id"] for s in api.get("/api/institute/backup/snapshots/").json()] == [snap.pk]
    assert api.get(f"/api/institute/backup/snapshots/{snap.pk}/").status_code == 404
    assert api.post(f"/api/institute/backup/snapshots/{snap.pk}/", {"confirm_code": "alpha"}, format="json").json()["restored"]
    assert set(Student.objects.filter(tenant=school_a).values_list("first_name", flat=True)) == {"Late"}

    # Teachers cannot restore; the owner can still use the same file.
    api.force_authenticate(teacher)
    assert api.post(url, {"file": upload(file)}, format="multipart").status_code == 403
    api.force_authenticate(super_admin)
    assert api.post(f"/api/platform/tenants/{school_a.pk}/restore/", {"file": upload(file)}, format="multipart").json()["signed"]


def test_school_admin_cannot_restore_changed_or_foreign_files(api, admin_a, school_a, make_tenant, make_user, tmp_path):
    beta = make_tenant("beta")
    api.force_authenticate(admin_a)
    url = "/api/institute/backup/restore/"
    good = json.loads(gzip.decompress(api.get("/api/institute/backup/").content))

    # A changed file (e.g. rows pointed at another institute) is refused.
    tampered = dict(good, rows=good["rows"] + [{"model": "students.student", "pk": 999,
                                                 "fields": {"tenant": beta.pk, "first_name": "X", "gender": "male"}}])
    res = api.post(url, {"file": upload(gzip.compress(json.dumps(tampered).encode()))}, format="multipart")
    assert res.status_code == 400 and "was changed" in str(res.json())
    unsigned = {k: v for k, v in good.items() if k != "signature"}
    assert "was changed" in str(api.post(url, {"file": upload(json.dumps(unsigned).encode())}, format="multipart").json())
    # Another institute's file, and a full platform backup (owner only).
    assert "for institute “beta”" in str(api.post(url, {"file": upload(institute.export_file(beta, False))}, format="multipart").json())
    with open(create_backup(str(tmp_path)), "rb") as full:
        assert "was changed" in str(api.post(url, {"file": upload(full.read())}, format="multipart").json())

    # A backup made before this admin's account would lock them out.
    old = institute.export_file(school_a, False)
    newcomer = make_user("new_admin", Role.ADMIN, school_a)
    api.force_authenticate(newcomer)
    assert "lock you out" in str(api.post(url, {"file": upload(old)}, format="multipart").json())
