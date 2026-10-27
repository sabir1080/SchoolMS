import datetime
import io

import pytest
from openpyxl import load_workbook

from apps.academics.models import AcademicSession, Program, Section, Subject
from apps.accounts.models import Role
from apps.attendance.models import AttendanceRecord
from apps.core.timeutils import tenant_today
from apps.exams.models import ExamSubject, Mark
from apps.fees.services import create_invoice, record_payment
from apps.payroll.models import Employee, Payslip
from apps.students.models import Enrollment, Guardian, Student, StudentGuardian


def auth(api, user):
    api.force_authenticate(user)
    return api


@pytest.fixture
def rp(school_a):
    today = tenant_today(school_a)
    session = AcademicSession.objects.create(tenant=school_a, name="Now", start_date=today - datetime.timedelta(days=100),
                                             end_date=today + datetime.timedelta(days=200), is_current=True)
    c5 = Program.objects.create(tenant=school_a, name="Class 5")
    c5.subjects.set([Subject.objects.create(tenant=school_a, name="English")])
    sec = Section.objects.create(tenant=school_a, program=c5, session=session, name="A")
    kids = []
    for roll, name in enumerate(["Zara", "Ali"], start=1):
        kid = Student.objects.create(tenant=school_a, admission_no=str(roll), first_name=name, gender="female")
        Enrollment.objects.create(tenant=school_a, student=kid, section=sec, session=session, roll_number=str(roll))
        StudentGuardian.objects.create(student=kid, guardian=Guardian.objects.create(tenant=school_a, name=f"{name} Father", phone="0300"), is_primary=True)
        kids.append(kid)
    return {"today": today, "session": session, "c5": c5, "sec": sec, "kids": kids}


def test_students_report_json_and_excel(api, admin_a, rp):
    api = auth(api, admin_a)
    body = api.get(f"/api/reports/students/?section={rp['sec'].pk}").json()
    assert body["title"] == "Students · Class 5 - A"
    assert [c["label"] for c in body["columns"]][:4] == ["Admission no", "Name", "Father / guardian", "Phone"]
    assert [r[1] for r in body["rows"]] == ["Zara", "Ali"] and body["totals"] == ["2 students"]
    res = api.get(f"/api/reports/students/?section={rp['sec'].pk}&download=xlsx")
    assert res.status_code == 200 and res["Content-Type"].startswith("application/vnd.openxmlformats")
    assert res["Content-Disposition"] == 'attachment; filename="students-class-5-a.xlsx"'
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws["A2"].value == "Students · Class 5 - A"
    assert [c.value for c in ws[4]][:2] == ["Admission no", "Name"]
    assert ws["B5"].value == "Zara" and ws["C5"].value == "Zara Father"


def test_attendance_fees_exam_payroll_reports(api, admin_a, rp, school_a):
    zara, ali = rp["kids"]
    today = rp["today"]
    for status in ("present", "absent"):
        AttendanceRecord.objects.create(tenant=school_a, student=zara, section=rp["sec"], status=status,
                                        date=today.replace(day=1) + datetime.timedelta(days=0 if status == "present" else 1))
    inv = create_invoice(school_a, ali, rp["session"], "September fee", today - datetime.timedelta(days=40), today - datetime.timedelta(days=30),
                         [{"description": "Tuition", "amount": 4000, "discount": 0}])
    record_payment(inv, 1500, today, "cash")
    api = auth(api, admin_a)
    att = api.get(f"/api/reports/attendance/?section={rp['sec'].pk}&month={today:%Y-%m}").json()
    assert att["rows"][0][:8] == ["1", "Zara", 2, 1, 0, 1, 0, 50.0]
    dues = api.get("/api/reports/fee-dues/").json()
    assert dues["rows"][0][1] == "Ali" and dues["rows"][0][6:9] == ["4000", "1500", "2500"] and dues["rows"][0][9] == 30
    assert dues["totals"][1] == "Total due: PKR 2,500"
    coll = api.get(f"/api/reports/fee-collection/?month={today:%Y-%m}").json()
    assert coll["rows"][0][-1] == "1500" and coll["totals"][1] == "Received: PKR 1,500"
    exam = api.post("/api/exams/", {"session": rp["session"].pk, "name": "First Term", "programs": [rp["c5"].pk]}, format="json").json()
    paper = ExamSubject.objects.get(exam_id=exam["id"])
    Mark.objects.create(tenant=school_a, paper=paper, student=zara, section=rp["sec"], marks_obtained="75.5")
    Mark.objects.create(tenant=school_a, paper=paper, student=ali, section=rp["sec"], is_absent=True)
    res = api.get(f"/api/reports/exam-result/?exam={exam['id']}&section={rp['sec'].pk}").json()
    assert res["columns"][3]["label"] == "English /100"
    assert res["rows"][0][:4] == [1, "1", "Zara", "75.5"] and res["rows"][1][3] == "Absent" and res["rows"][1][-1] == "Fail"
    sana = Employee.objects.create(tenant=school_a, name="Sana", basic_salary=40000)
    Payslip.objects.create(tenant=school_a, employee=sana, period=f"{today:%Y-%m}", basic=40000, deduction=1000)
    pay = api.get(f"/api/reports/payroll/?month={today:%Y-%m}").json()
    assert pay["rows"][0][6] == "39000" and pay["totals"][0] == "Total: PKR 39,000"
    assert api.get(f"/api/reports/payroll/?month={today:%Y-%m}&download=xlsx").status_code == 200


def test_report_access_and_plans(api, admin_a, admin_b, make_user, rp, school_a):
    api = auth(api, admin_a)
    assert api.get("/api/reports/nope/").status_code == 404
    assert api.get("/api/reports/attendance/").status_code == 400  # section required
    assert api.get("/api/reports/exam-result/?exam=999").status_code == 400
    assert auth(api, make_user("t", Role.TEACHER, school_a)).get("/api/reports/students/").status_code == 403
    assert auth(api, admin_b).get(f"/api/reports/students/?section={rp['sec'].pk}").status_code == 400  # other institute's section
    plan = school_a.subscription.plan
    plan.features = [f for f in plan.features if f != "payroll"]
    plan.save()
    assert auth(api, admin_a).get("/api/reports/payroll/").status_code == 402
    plan.features = [f for f in plan.features if f != "reports"]
    plan.save()
    assert auth(api, admin_a).get("/api/reports/students/").status_code == 402
