"""An institute's main records as an Excel workbook people can read (students, bills, payments, results,
attendance and staff). It is a copy for the institute's own files; restoring uses the backup file instead."""

import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.academics.models import Section
from apps.attendance.models import AttendanceRecord, summarize
from apps.exams.models import Exam, ExamStatus
from apps.exams.results import section_results
from apps.fees.models import Invoice, Payment
from apps.payroll.models import Employee
from apps.students.models import Student
from apps.students.serializers import current_enrollment, primary_link


def _sheet(wb, title, columns, rows, institute):
    ws = wb.create_sheet(title)
    ws.append([institute])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([title])
    ws["A2"].font = Font(bold=True, color="4F46E5")
    ws.append([])
    ws.append(columns)
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="EEF2FF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for index, label in enumerate(columns, start=1):
        width = max([len(str(label))] + [len(str(r[index - 1])) for r in rows if r[index - 1] is not None] or [8])
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A5"


def _students(tenant):
    students = Student.objects.filter(tenant=tenant).prefetch_related(
        "enrollments__section__program", "enrollments__session", "guardian_links__guardian").order_by("first_name", "last_name")
    rows = []
    for s in students:
        e = current_enrollment(s)
        link = primary_link(s)
        rows.append([s.admission_no, s.full_name, s.get_gender_display(), s.date_of_birth,
                     f"{e.section.program.name} - {e.section.name}" if e else "", e.roll_number if e else "",
                     s.get_status_display(), link.guardian.name if link else "", link.guardian.phone if link else "", s.admission_date])
    return ["Admission no", "Name", "Gender", "Date of birth", "Class", "Roll no", "Status", "Parent / guardian", "Phone", "Admission date"], rows


def _bills(tenant):
    rows = [[i.number, i.student.full_name, i.student.admission_no, i.title, i.period, i.issue_date, i.due_date,
             i.total, i.paid, i.total - i.paid, i.get_status_display()]
            for i in Invoice.objects.filter(tenant=tenant).select_related("student").order_by("number")]
    return ["Bill no", "Student", "Admission no", "Title", "Month", "Issued", "Due", "Total", "Paid", "Balance", "Status"], rows


def _payments(tenant):
    rows = [[p.receipt_number, p.date, p.invoice.student.full_name, p.invoice.title, p.amount, p.get_method_display(), p.reference]
            for p in Payment.objects.filter(tenant=tenant).select_related("invoice__student").order_by("receipt_number")]
    return ["Receipt no", "Date", "Student", "Bill", "Amount", "Method", "Reference"], rows


def _results(tenant):
    rows = []
    for exam in Exam.objects.filter(tenant=tenant, status=ExamStatus.PUBLISHED).order_by("start_date", "pk"):
        for section in Section.objects.filter(tenant=tenant, session=exam.session, program__in=exam.programs.all()).select_related("program"):
            for r in section_results(exam, section)["students"]:
                rows.append([exam.name, f"{section.program.name} - {section.name}", r.get("roll_number", ""), r.get("full_name"),
                             r.get("percentage"), r.get("grade"), (r.get("result") or "").capitalize(), r.get("position")])
    return ["Exam", "Class", "Roll no", "Student", "Percentage", "Grade", "Result", "Position"], rows


def _attendance(tenant):
    statuses = defaultdict(list)
    names = {}
    for rec in AttendanceRecord.objects.filter(tenant=tenant, section__session__is_current=True).select_related("student").order_by():
        statuses[rec.student_id].append(rec.status)
        names[rec.student_id] = (rec.student.admission_no, rec.student.full_name)
    rows = []
    for sid, (adm, name) in sorted(names.items(), key=lambda x: x[1][1]):
        s = summarize(statuses[sid])
        rows.append([adm, name, s["present"], s["absent"], s["leave"], s["late"], s["total"], s["percentage"]])
    return ["Admission no", "Student", "Present", "Absent", "Leave", "Late", "Days", "Attendance %"], rows


def _staff(tenant):
    rows = [[e.name, e.designation, e.phone, e.join_date, e.basic_salary, e.allowance, "Active" if e.is_active else "Left"]
            for e in Employee.objects.filter(tenant=tenant).order_by("name")]
    return ["Name", "Designation", "Phone", "Joining date", "Basic salary", "Allowance", "Status"], rows


SHEETS = [("Students", _students), ("Bills", _bills), ("Payments", _payments), ("Results", _results),
          ("Attendance (this session)", _attendance), ("Staff", _staff)]


def institute_workbook(tenant):
    wb = Workbook()
    wb.remove(wb.active)
    for title, build in SHEETS:
        columns, rows = build(tenant)
        _sheet(wb, title, columns, rows, tenant.name)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
