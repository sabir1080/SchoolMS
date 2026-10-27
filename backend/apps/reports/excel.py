import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATS = {"money": "#,##0", "number": "0", "percent": "0.0", "date": "dd-mmm-yyyy"}


def workbook(tenant, report):
    """An .xlsx file: institute name, report title, a bold header row, the rows and the totals."""
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:31].replace("/", "-").replace(":", "-") or "Report"
    columns = report["columns"]
    ws.append([tenant.name])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([report["title"]])
    ws["A2"].font = Font(bold=True, size=12, color="4F46E5")
    ws.append([])
    ws.append([label for label, _kind in columns])
    header_fill = PatternFill("solid", fgColor="EEF2FF")
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True), header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in report["rows"]:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        for index, (_label, kind) in enumerate(columns, start=1):
            if kind in FORMATS:
                ws.cell(row=ws.max_row, column=index).number_format = FORMATS[kind]
    if report["totals"]:
        ws.append([])
        for line in report["totals"]:
            ws.append([line])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for index, (label, _kind) in enumerate(columns, start=1):
        values = [label] + [row[index - 1] for row in report["rows"]]
        width = max(len(str(v)) if v is not None else 0 for v in values)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 40)
    ws.freeze_panes = "A5"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
